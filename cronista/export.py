"""Transcript files: SRT and plain text next to the video, and Word or PDF on demand.

write_word is escriba's DocumentService.generate (backend/app/services/document.py) with its
headings in English or Spanish (language.py) and without the language-detection line: the full text in readable paragraphs first,
then every segment with its time range. write_pdf lays out the same with reportlab, in a Windows
TrueType font embedded in the file (Segoe UI, else Arial), so accents, ñ and ¿¡ show anywhere.

With `names` (voices.speakers: the label of each audio source, when both spoke), every file says
who said what: the reading part goes by turns ("Ana: ..."), and each timed line has its speaker.
"""
from __future__ import annotations

import itertools
import os
import threading
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .engine.transcript import TranscriptionSegment, to_srt
from .engine.voices import line
from .language import export_words

# (regular, bold), in order of preference; reportlab's built-in Helvetica if none is there.
PDF_FONT_FILES = [
    (r"%WINDIR%\Fonts\segoeui.ttf", r"%WINDIR%\Fonts\segoeuib.ttf"),
    (r"%WINDIR%\Fonts\arial.ttf", r"%WINDIR%\Fonts\arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]
_pdf_fonts_lock = threading.Lock()
_pdf_fonts_found: tuple[str, str] | None = None


def _clock(seconds: float) -> str:
    hours, rest = divmod(int(seconds), 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _time_range(segment: TranscriptionSegment) -> str:
    return f"[{_clock(segment.start)} → {_clock(segment.end)}]"


def _reading_paragraphs(segments: list[TranscriptionSegment]) -> list[str]:
    """The whole text in paragraphs of ~500 characters, each ending at a sentence boundary."""
    paragraphs: list[str] = []
    buffer: list[str] = []
    length = 0
    for word in " ".join(segment.text for segment in segments).split():
        buffer.append(word)
        length += len(word) + 1
        if length > 500 and word.endswith((".", "?", "!", ":")):
            paragraphs.append(" ".join(buffer))
            buffer, length = [], 0
    if buffer:
        paragraphs.append(" ".join(buffer))
    return paragraphs


def _turns(segments: list[TranscriptionSegment], names: dict[str, str] | None) -> list[tuple[str | None, str]]:
    """The paragraphs to read, as (speaker, text): with names, a new one each time the speaker changes
    (a long turn is split like _reading_paragraphs, the speaker said once); without, as before."""
    if not names:
        return [(None, paragraph) for paragraph in _reading_paragraphs(segments)]
    turns: list[tuple[str | None, str]] = []
    for source, said in itertools.groupby(segments, key=lambda segment: segment.source):
        for index, paragraph in enumerate(_reading_paragraphs(list(said))):
            turns.append((names.get(source) if index == 0 else None, paragraph))
    return turns


def write_text_files(segments: list[TranscriptionSegment], folder: Path,
                     names: dict[str, str] | None = None) -> tuple[Path, Path]:
    folder = Path(folder)
    srt = folder / "transcript.srt"
    txt = folder / "transcript.txt"
    srt.write_text(to_srt(segments, names), encoding="utf-8")
    txt.write_text("".join(f"{line(segment, names)}\n" for segment in segments), encoding="utf-8")
    return srt, txt


def write_word(segments: list[TranscriptionSegment], output_path: Path, *, title: str, recorded_at: datetime,
               language: str = "en", names: dict[str, str] | None = None) -> Path:
    words = export_words(language)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(words["recorded"].format(recorded_at))
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_heading(words["transcript"], level=2)
    for speaker, text in _turns(segments, names):
        paragraph = doc.add_paragraph()
        if speaker:
            paragraph.add_run(f"{speaker}: ").bold = True
        paragraph.add_run(text)

    doc.add_heading(words["timestamps"], level=2)
    for segment in segments:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        stamp = paragraph.add_run(f"{_time_range(segment)} ")
        stamp.bold = True
        stamp.font.size = Pt(9)
        stamp.font.color.rgb = RGBColor(70, 70, 70)
        if names and segment.source in names:
            speaker = paragraph.add_run(f"{names[segment.source]}: ")
            speaker.bold = True
            speaker.font.size = Pt(10)
        paragraph.add_run(segment.text).font.size = Pt(10)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output


def _pdf_fonts() -> tuple[str, str]:
    """The PDF's (regular, bold) font names, registered with reportlab the first time."""
    global _pdf_fonts_found
    with _pdf_fonts_lock:
        if _pdf_fonts_found is None:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont

            _pdf_fonts_found = ("Helvetica", "Helvetica-Bold")  # built in: accents fine, rarer characters not
            for regular, bold in PDF_FONT_FILES:
                regular, bold = os.path.expandvars(regular), os.path.expandvars(bold)
                if os.path.isfile(regular) and os.path.isfile(bold):
                    pdfmetrics.registerFont(TTFont("Transcript", regular))
                    pdfmetrics.registerFont(TTFont("Transcript-Bold", bold))
                    pdfmetrics.registerFontFamily("Transcript", normal="Transcript", bold="Transcript-Bold",
                                                  italic="Transcript", boldItalic="Transcript-Bold")
                    _pdf_fonts_found = ("Transcript", "Transcript-Bold")
                    break
        return _pdf_fonts_found


def pdf_story(segments: list[TranscriptionSegment], *, title: str, recorded_at: datetime, language: str = "en",
              names: dict[str, str] | None = None) -> list:
    """The PDF's content as reportlab flowables: what write_word writes, in the same order."""
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph

    words = export_words(language)
    regular, bold = _pdf_fonts()
    title_style = ParagraphStyle("title", fontName=bold, fontSize=18, leading=23, alignment=TA_CENTER, spaceAfter=4)
    meta_style = ParagraphStyle("meta", fontName=regular, fontSize=10, leading=13, alignment=TA_CENTER,
                                textColor=HexColor("#646464"), spaceAfter=16)
    heading = ParagraphStyle("heading", fontName=bold, fontSize=13, leading=17, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("body", fontName=regular, fontSize=11, leading=15.5, spaceAfter=7)
    timed_style = ParagraphStyle("line", fontName=regular, fontSize=10, leading=13.5, spaceAfter=2.5)

    story = [Paragraph(escape(title), title_style),
             Paragraph(escape(words["recorded"].format(recorded_at)), meta_style),
             Paragraph(words["transcript"], heading)]
    story += [Paragraph((f"<b>{escape(speaker)}:</b> " if speaker else "") + escape(text), body)
              for speaker, text in _turns(segments, names)]
    story.append(Paragraph(words["timestamps"], heading))

    def timed(segment: TranscriptionSegment) -> str:
        speaker = f"<b>{escape(names[segment.source])}:</b> " if names and segment.source in names else ""
        return f'<font size="9" color="#464646"><b>{_time_range(segment)}</b></font> {speaker}{escape(segment.text)}'

    story += [Paragraph(timed(segment), timed_style) for segment in segments]
    return story


def write_pdf(segments: list[TranscriptionSegment], output_path: Path, *, title: str, recorded_at: datetime,
              language: str = "en", names: dict[str, str] | None = None) -> Path:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate

    regular, _ = _pdf_fonts()
    margin = 2.2 * cm

    def page_number(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(HexColor("#8a8a8a"))
        canvas.drawRightString(A4[0] - margin, 1.2 * cm, str(document.page))
        canvas.restoreState()

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(str(output), pagesize=A4, title=title, author="Cronista",
                                 leftMargin=margin, rightMargin=margin, topMargin=2 * cm, bottomMargin=2 * cm)
    document.build(pdf_story(segments, title=title, recorded_at=recorded_at, language=language, names=names),
                   onFirstPage=page_number, onLaterPages=page_number)
    return output
