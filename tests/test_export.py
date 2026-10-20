import os
import re
from datetime import datetime

from docx import Document

from cronista.engine.transcript import TranscriptionSegment
from cronista.export import pdf_story, write_pdf, write_text_files, write_word

SEGMENTS = [TranscriptionSegment(0.0, 2.5, "Hola."), TranscriptionSegment(3.0, 65.0, "Second part.")]
RECORDED = datetime(2026, 10, 19, 11, 20)


def test_word_has_the_full_text_and_a_timestamped_section(tmp_path):
    path = write_word(SEGMENTS, tmp_path / "transcript.docx", title="Demo", recorded_at=RECORDED)
    paragraphs = [p.text for p in Document(str(path)).paragraphs]
    assert paragraphs[0] == "Demo"
    assert "Hola. Second part." in paragraphs
    assert "[00:00:00 → 00:00:02] Hola." in paragraphs
    assert "[00:00:03 → 00:01:05] Second part." in paragraphs


def test_pdf_says_what_the_word_file_says():
    texts = [flowable.getPlainText() for flowable in pdf_story(SEGMENTS, title="Demo & <co>", recorded_at=RECORDED)
             if hasattr(flowable, "getPlainText")]
    assert texts[0] == "Demo & <co>"  # escaped for reportlab's markup, shown as typed
    assert texts[1] == "Recorded on 2026-10-19 11:20"
    assert "Hola. Second part." in texts
    assert "[00:00:00 → 00:00:02] Hola." in texts
    assert "[00:00:03 → 00:01:05] Second part." in texts


def test_word_and_pdf_in_spanish(tmp_path):
    path = write_word(SEGMENTS, tmp_path / "transcript.docx", title="Demo", recorded_at=RECORDED, language="es")
    paragraphs = [p.text for p in Document(str(path)).paragraphs]
    assert paragraphs[1:3] == ["Grabado el 19/10/2026 a las 11:20", "Transcripción"]
    assert "Marcas de tiempo" in paragraphs
    texts = [flowable.getPlainText() for flowable in pdf_story(SEGMENTS, title="Demo", recorded_at=RECORDED, language="es")
             if hasattr(flowable, "getPlainText")]
    assert texts[1:3] == ["Grabado el 19/10/2026 a las 11:20", "Transcripción"]
    assert "Marcas de tiempo" in texts


def test_pdf_is_a_valid_file_with_its_font_inside_and_room_for_long_transcripts(tmp_path):
    sentence = "Una frase con acentos, eñes y signos: ¿así? ¡Sí! Ünicode → listo."
    long = [TranscriptionSegment(i * 5.0, i * 5.0 + 4.0, sentence) for i in range(300)]
    path = write_pdf(long, tmp_path / "transcript.pdf", title="Grabación larga", recorded_at=RECORDED)
    data = path.read_bytes()
    assert data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF")
    assert len(re.findall(rb"/Type /Page[^s]", data)) > 5
    if os.name == "nt":
        assert b"/FontFile2" in data  # a Windows TrueType font, embedded: accents show anywhere


def test_writes_srt_and_plain_text(tmp_path):
    srt, txt = write_text_files(SEGMENTS, tmp_path)
    assert srt.read_text(encoding="utf-8").startswith("1\n00:00:00,000 --> 00:00:02,500\nHola.\n")
    assert txt.read_text(encoding="utf-8") == "Hola.\nSecond part.\n"


CALL = [TranscriptionSegment(0.0, 2.0, "Let's start.", "mic"),
        TranscriptionSegment(2.5, 4.0, "Sure, the screens are ready.", "computer"),
        TranscriptionSegment(4.2, 5.0, "Show them.", "computer"),
        TranscriptionSegment(6.0, 7.0, "Great.", "mic")]
NAMES = {"mic": "Me", "computer": "Ana"}


def test_with_two_sources_every_file_says_who_said_what(tmp_path):
    srt, txt = write_text_files(CALL, tmp_path, NAMES)
    assert "Ana: Sure, the screens are ready." in srt.read_text(encoding="utf-8")
    assert txt.read_text(encoding="utf-8").splitlines() == [
        "Me: Let's start.", "Ana: Sure, the screens are ready.", "Ana: Show them.", "Me: Great."]

    path = write_word(CALL, tmp_path / "transcript.docx", title="Call", recorded_at=RECORDED, names=NAMES)
    paragraphs = [p.text for p in Document(str(path)).paragraphs]
    # The reading part goes by turns: one paragraph each time the speaker changes.
    assert "Ana: Sure, the screens are ready. Show them." in paragraphs
    assert "Me: Great." in paragraphs
    assert "[00:00:02 → 00:00:04] Ana: Sure, the screens are ready." in paragraphs

    texts = [flowable.getPlainText() for flowable in pdf_story(CALL, title="Call", recorded_at=RECORDED, names=NAMES)
             if hasattr(flowable, "getPlainText")]
    assert "Ana: Sure, the screens are ready. Show them." in texts
    assert "[00:00:02 → 00:00:04] Ana: Sure, the screens are ready." in texts
