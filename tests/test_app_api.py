"""The window's API on its own: a history in a temporary folder, no recorder and no model."""
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

import cronista.app_api as app_api
from cronista import importer
from cronista.downloads import DownloadCancelled, DownloadFailed, RemoteFile
from cronista.engine.canary_pass import CanaryPass
from cronista.engine.transcript import TranscriptionSegment


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(app_api, "data_dir", lambda: tmp_path / "data")
    monkeypatch.setattr(app_api, "recordings_dir", lambda: tmp_path / "videos")
    monkeypatch.setattr(app_api, "IS_WINDOWS", False)  # no Explorer windows from the tests
    monkeypatch.setattr(app_api, "default_recorder", lambda: None)
    return app_api.Api()


def add_call(api, tmp_path):
    folder = tmp_path / "videos" / "2026-10-20_10-00-00"
    folder.mkdir(parents=True)
    segments = [{"start": 0.0, "end": 1.0, "text": "Let's start.", "source": "mic"},
                {"start": 1.5, "end": 3.0, "text": "Sure.", "source": "computer"}]
    recording_id = api._library.add(
        title="Call", created_at="2026-10-20T10:00:00", duration_seconds=3.0, size_bytes=1, folder=str(folder),
        video_path=str(folder / "recording.mp4"), transcript="Let's start. Sure.", segments=segments,
        labels={"mic": "Me", "computer": "Others"})
    return recording_id, folder


def test_renaming_a_label_renames_it_in_the_files_too(api, tmp_path):
    recording_id, folder = add_call(api, tmp_path)
    assert api.rename_label(recording_id, "computer", "  Ana  ") == {"ok": True, "labels": {"mic": "Me", "computer": "Ana"}}
    assert "Ana: Sure." in (folder / "transcript.srt").read_text(encoding="utf-8")
    assert (folder / "transcript.txt").read_text(encoding="utf-8").splitlines() == ["Me: Let's start.", "Ana: Sure."]
    assert api.rename_label(recording_id, "computer", "   ") == {"error": "Enter a label."}
    assert api.rename_label(recording_id, "screen", "X") == {"error": "Recording not found."}


def test_exports_say_who_said_what(api, tmp_path):
    recording_id, _ = add_call(api, tmp_path)
    path = api.export_word(recording_id, "en")["path"]
    assert "[00:00:01 → 00:00:03] Others: Sure." in [paragraph.text for paragraph in Document(path).paragraphs]


def test_the_default_labels_are_kept_only_when_named(api):
    assert api.state()["settings"]["labels"] == {}
    api.save_preferences(labels={"mic": " Nacho ", "computer": "", "screen": "x"})
    assert api.state()["settings"]["labels"] == {"mic": "Nacho"}
    api.save_preferences(subtitles=True)
    assert api.state()["settings"]["subtitles"] is True
    assert api.state()["settings"]["labels"] == {"mic": "Nacho"}  # one preference does not undo another


def test_the_subtitled_video_is_made_in_the_background_with_the_labels(api, tmp_path, monkeypatch):
    import threading
    from types import SimpleNamespace

    recording_id, folder = add_call(api, tmp_path)
    (folder / "recording.mp4").write_bytes(b"video")
    api.rename_label(recording_id, "computer", "Ana")
    made = []

    def fake_burn(ffmpeg, where, segments, names, *, progress):
        made.append((ffmpeg, where, [s.source for s in segments], names))
        for seconds in (1.5, 2.25, 3.2):  # the call is 3 s long; ffmpeg's last word can go past that
            progress(seconds)
        (where / "recording-subtitled.mp4").write_bytes(b"video")
        return where / "recording-subtitled.mp4"

    monkeypatch.setattr(app_api, "burn_subtitles", fake_burn)
    api._recorder = SimpleNamespace(ffmpeg="ffmpeg.exe")
    events, done = [], threading.Event()
    api._emit = lambda kind, payload: (events.append((kind, payload)), kind == "subtitled" and done.set())
    assert api.subtitled_video(recording_id, "en") == {"ok": True}
    assert done.wait(5)
    assert made == [("ffmpeg.exe", folder, ["mic", "computer"], {"mic": "Me", "computer": "Ana"})]
    # How much is done, never all of it before "subtitled", and the seconds left at the pace so far
    # (none here: the fake is instant).
    assert [payload for kind, payload in events if kind == "subtitling"] == [
        {"id": recording_id, "progress": 0.0, "eta": None},
        {"id": recording_id, "progress": 0.5, "eta": 0},
        {"id": recording_id, "progress": 0.75, "eta": 0},
        {"id": recording_id, "progress": 0.99, "eta": 0}]
    assert events[-1] == ("subtitled", {"id": recording_id, "path": str(folder / "recording-subtitled.mp4")})


def test_copied_text_says_who_said_what(api, tmp_path):
    recording_id, _ = add_call(api, tmp_path)
    assert api.copy_text(recording_id) == {"text": "Me: Let's start.\nOthers: Sure."}


def test_the_theme_is_one_of_the_known_ones(api):
    assert api.state()["settings"]["theme"] == "system"
    api.save_preferences(theme="midnight")
    assert api.state()["settings"]["theme"] == "midnight"
    api.save_preferences(theme="neon-pink")  # not one of the page's themes: left as it was
    assert api.state()["settings"]["theme"] == "midnight"


# --- transcribing a file the user drops or picks ------------------------------------------------

def a_file(tmp_path, name="Interview with Laura.m4a", size=1000):
    path = tmp_path / "Downloads" / name
    path.parent.mkdir(exist_ok=True)
    path.write_bytes(b"x" * size)
    return path


def fake_transcribe(said="Thanks for coming.", *, has_video=False, wait=None):
    """Stands in for importer.transcribe_file: reports its progress and says `said`, or waits for a stop."""

    def transcribe(model, ffmpeg, path, folder, *, progress, lines, stopped, refine=None):
        progress("reading", 0.0, 12.0)
        if wait is not None:
            while not stopped():
                wait.wait(0.01)
            raise importer.Cancelled
        progress("transcribing", 6.0, 12.0)
        segments = [TranscriptionSegment(0.0, 2.0, said)]
        lines(segments)
        progress("transcribing", 12.0, 12.0)
        return importer.ImportResult(duration=12.0, has_video=has_video, segments=segments)

    return transcribe


@pytest.fixture
def importing(api, monkeypatch):
    """The API ready to transcribe files, its events in a list; `until(kind)` waits for one."""
    import threading

    monkeypatch.setattr(app_api, "find_ffmpeg", lambda: "ffmpeg.exe")
    api._model = object()
    events, arrived = [], threading.Condition()

    def emit(kind, payload):
        with arrived:
            events.append((kind, payload))
            arrived.notify_all()

    def until(kind, count=1):
        with arrived:
            assert arrived.wait_for(lambda: sum(k == kind for k, _ in events) >= count, timeout=5), events
        return [payload for k, payload in events if k == kind]

    api._emit = emit
    return SimpleNamespace(events=events, until=until)


def test_a_dropped_file_is_transcribed_into_the_history_and_the_original_is_kept(api, tmp_path, monkeypatch,
                                                                                 importing):
    monkeypatch.setattr(app_api, "transcribe_file", fake_transcribe())
    original = a_file(tmp_path)
    assert api._import_files([str(original)]) == {"ok": True, "queued": 1}
    recording_id = importing.until("imported")[0]["id"]
    record = api.recording(recording_id)
    assert record["title"] == "Interview with Laura" and record["kind"] == "audio"
    assert record["video_path"] == str(original)  # Play and Show in folder open the file itself
    assert record["duration_seconds"] == 12.0 and record["size_bytes"] == 1000
    assert (Path(record["folder"]) / "transcript.txt").read_text(encoding="utf-8") == "Thanks for coming.\n"
    progress = importing.until("import")
    assert progress[0]["name"] == "Interview with Laura.m4a" and progress[0]["stage"] == "reading"
    assert any(p.get("lines") == [{"start": 0.0, "end": 2.0, "text": "Thanks for coming."}] for p in progress)
    assert progress[-1]["progress"] == 1.0
    assert api.state()["importing"] is None

    assert api.delete(recording_id) == {"ok": True}
    assert original.exists()  # deleting it from the history never deletes the user's file
    assert not Path(record["folder"]).exists()


def test_files_dropped_together_are_transcribed_one_after_another(api, tmp_path, monkeypatch, importing):
    monkeypatch.setattr(app_api, "transcribe_file", fake_transcribe())
    first, second = a_file(tmp_path, "one.mp3"), a_file(tmp_path, "two.mkv")
    folder = tmp_path / "Downloads" / "a folder"
    folder.mkdir()
    assert api._import_files([str(first), str(folder), str(second)]) == {"ok": True, "queued": 2}  # files only
    importing.until("imported", 2)
    assert [r["title"] for r in api.recordings()] == ["two", "one"]
    assert importing.until("import")[0]["waiting"] == 1


def test_one_thing_at_a_time_a_recording_or_a_file(api, tmp_path, monkeypatch, importing):
    import threading

    api._phase = "recording"
    assert api._import_files([str(a_file(tmp_path))]) == {"error": "Stop the recording first, then drop the file again."}
    api._phase = "idle"
    monkeypatch.setattr(app_api, "transcribe_file", fake_transcribe(wait=threading.Event()))
    assert api._import_files([str(a_file(tmp_path))])["ok"]
    importing.until("import")
    api._recorder = SimpleNamespace()
    assert api.start(None, "balanced") == {"error": "A file is being transcribed. Cancel it to record."}
    assert api.state()["importing"]["name"] == "Interview with Laura.m4a"
    assert api.cancel_import() == {"ok": True}
    importing.until("import-cancelled")
    assert api.recordings() == [] and list((tmp_path / "videos").iterdir()) == []  # nothing half-done stays


def test_a_file_that_cannot_be_transcribed_says_why(api, tmp_path, monkeypatch, importing):
    def failing(*_, **__):
        raise importer.ImportFailed("This file has no sound to transcribe.")

    monkeypatch.setattr(app_api, "transcribe_file", failing)
    api._import_files([str(a_file(tmp_path, "screen.mp4"))])
    error = importing.until("import-error")[0]
    assert error == {"name": "screen.mp4", "message": "This file has no sound to transcribe."}
    assert api.recordings() == [] and list((tmp_path / "videos").iterdir()) == []


def test_a_drop_brings_the_files_full_paths(api, importing, monkeypatch):
    queued = []
    monkeypatch.setattr(api, "_import_files", lambda paths: queued.append(paths) or {"ok": True, "queued": len(paths)})
    api._on_drop({"dataTransfer": {"files": [{"name": "a.mp3", "pywebviewFullPath": r"C:\talks\a.mp3"},
                                             {"name": "b.txt"}]}})
    assert queued == [[r"C:\talks\a.mp3"]]
    api._on_drop({"dataTransfer": {"files": [{"name": "b.txt"}]}})
    assert importing.until("import-refused")[0]["message"].startswith("The dropped file could not be read.")


def test_a_drop_while_recording_is_refused_without_ending_the_recording(api, tmp_path, importing):
    api._phase = "recording"
    api._on_drop({"dataTransfer": {"files": [{"name": "a.mp3", "pywebviewFullPath": str(a_file(tmp_path))}]}})
    assert importing.until("import-refused")[0]["message"] == "Stop the recording first, then drop the file again."
    assert not any(kind == "error" for kind, _ in importing.events)  # the page ends the recording on "error"


def test_a_dropped_file_in_high_quality_is_written_again_by_canary(api, tmp_path, monkeypatch, importing,
                                                                     canary_files):
    canary_here(api, canary_files)
    assert api.set_transcription("high") == {"ok": True, "transcription": "high"}
    loaded = object()
    monkeypatch.setattr(app_api, "load_canary", lambda folder, threads: loaded)
    refines = []
    plain = fake_transcribe()

    def transcribe(model, ffmpeg, path, folder, *, progress, lines, stopped, refine=None):
        refines.append(refine)
        return plain(model, ffmpeg, path, folder, progress=progress, lines=lines, stopped=stopped)

    monkeypatch.setattr(app_api, "transcribe_file", transcribe)
    api._import_files([str(a_file(tmp_path))])
    importing.until("imported")
    assert isinstance(refines[0], CanaryPass) and refines[0].model is loaded
    wait_until(lambda: api._importer is None)
    assert api._canary is None  # about 1 GB of memory back once the files are done


def test_a_transcribed_files_documents_say_when_it_was_transcribed(api, tmp_path, monkeypatch, importing):
    monkeypatch.setattr(app_api, "transcribe_file", fake_transcribe(has_video=True))
    api._import_files([str(a_file(tmp_path, "Demo.mkv"))])
    recording_id = importing.until("imported")[0]["id"]
    assert api.recording(recording_id)["kind"] == "video"
    paragraphs = [paragraph.text for paragraph in Document(api.export_word(recording_id, "en")["path"]).paragraphs]
    assert any(text.startswith("Transcribed on ") for text in paragraphs)
    assert not any(text.startswith("Recorded on ") for text in paragraphs)


# --- transcription quality: medium-low (Parakeet) or high (Canary, downloaded the first time) -----

@pytest.fixture
def canary_files(monkeypatch):
    """Canary as three tiny files, so it downloads in an instant."""
    files = [RemoteFile(name, f"https://example.invalid/{name}", size, "-")
             for name, size in (("config.json", 3), ("vocab.txt", 5), ("encoder-model.int8.onnx", 8))]
    monkeypatch.setattr(app_api, "CANARY_FILES", files)
    return files


def canary_here(api, files):
    folder = api._canary_folder()
    folder.mkdir(parents=True, exist_ok=True)
    for file in files:
        (folder / file.name).write_bytes(b"x" * file.size)


def fake_download(*, wait=None, fail=None):
    """Stands in for downloads.download: writes the files, or waits for a stop, or fails."""

    def download(folder, files, *, progress, stopped):
        total = sum(file.size for file in files)
        progress(0, total)
        if wait is not None:
            while not stopped():
                wait.wait(0.01)
            raise DownloadCancelled
        if fail:
            raise DownloadFailed(fail)
        folder.mkdir(parents=True, exist_ok=True)
        done = 0
        for file in files:
            (folder / file.name).write_bytes(b"x" * file.size)
            done += file.size
            progress(done, total)

    return download


def wait_until(check, timeout=5.0):
    import time

    deadline = time.monotonic() + timeout
    while not check():
        assert time.monotonic() < deadline, "timed out"
        time.sleep(0.01)


def test_high_quality_downloads_canary_once_and_then_is_chosen(api, monkeypatch, importing, canary_files):
    monkeypatch.setattr(app_api, "download", fake_download())
    assert api.state()["canary"] == {"state": "missing"}
    assert api.state()["settings"]["transcription"] == "standard"
    assert api.set_transcription("high") == {"ok": True, "downloading": True}
    wait_until(lambda: api.state()["canary"]["state"] == "ready")
    events = [payload for kind, payload in importing.events if kind == "canary"]
    assert any(event["state"] == "downloading" and event["total"] == 16 for event in events)
    assert events[-1] == {"state": "ready", "transcription": "high"}
    assert api.state()["settings"]["transcription"] == "high"
    assert api.set_transcription("high") == {"ok": True, "transcription": "high"}  # here now: no download
    assert api.set_transcription("standard") == {"ok": True, "transcription": "standard"}
    assert api.state()["settings"]["transcription"] == "standard"


def test_choosing_medium_low_while_canary_downloads_stops_the_download(api, monkeypatch, importing, canary_files):
    import threading

    monkeypatch.setattr(app_api, "download", fake_download(wait=threading.Event()))
    api.set_transcription("high")
    wait_until(lambda: api.state()["canary"]["state"] == "downloading")
    api.set_transcription("standard")
    wait_until(lambda: api.state()["canary"]["state"] == "missing")
    assert [p for k, p in importing.events if k == "canary"][-1] == {"state": "missing", "transcription": "standard"}


def test_a_download_that_fails_says_why_and_stays_medium_low(api, monkeypatch, importing, canary_files):
    message = "Downloading the model failed (no network). Check the connection and try again."
    monkeypatch.setattr(app_api, "download", fake_download(fail=message))
    api.set_transcription("high")
    wait_until(lambda: [p for k, p in importing.events if k == "canary"][-1]["state"] != "downloading")
    assert [p for k, p in importing.events if k == "canary"][-1] == {"state": "error", "message": message,
                                                                    "transcription": "standard"}
    assert api.state()["canary"] == {"state": "missing"}


def test_a_recording_in_high_quality_gets_canary_for_each_source(api, monkeypatch, importing, canary_files):
    made = []

    class FakeSession:
        def __init__(self, model, recorder, settings, folder, on_event, *, refiner=None):
            made.append(refiner)

        def start(self):
            pass

    monkeypatch.setattr(app_api, "RecordingSession", FakeSession)
    api._recorder, api._model = SimpleNamespace(), object()
    assert api.start(None, "balanced")["ok"] and made == [None]  # medium-low: Parakeet alone
    api._phase = "idle"
    canary_here(api, canary_files)
    api.set_transcription("high")
    loaded = object()
    monkeypatch.setattr(app_api, "load_canary", lambda folder, threads: loaded)
    assert api.start(None, "balanced")["ok"]
    first, second = made[1](), made[1]()
    assert isinstance(first, CanaryPass) and first is not second and first.model is second.model is loaded


def test_if_canary_does_not_load_the_transcript_is_medium_low_and_the_page_is_told(api, monkeypatch, importing,
                                                                                    canary_files):
    canary_here(api, canary_files)
    api.set_transcription("high")

    def broken(folder, threads):
        raise RuntimeError("bad file")

    monkeypatch.setattr(app_api, "load_canary", broken)
    assert api._refiner()() is None
    assert importing.until("notice")[0]["message"].startswith("High quality could not be used (bad file).")
